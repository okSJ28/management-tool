"""Excel / CSV import and export."""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any, Iterable

from . import model

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:  # pragma: no cover
    HAS_OPENPYXL = False

MAX_BLANK_RUN = 60
KEY_HEADER = model.norm_header("プログラムID")


# ---------------------------------------------------------------- helpers
def _alias_map(type_key: str) -> dict[str, str]:
    aliases: dict[str, str] = {}
    for field in model.SCHEMA["fields"]:
        if type_key not in field["types"]:
            continue
        for alias in field.get("excel", []) or []:
            aliases[model.norm_header(alias)] = field["key"]
        aliases.setdefault(model.norm_header(field["label"]), field["key"])
    return aliases


def _match_key(header: str, aliases: dict[str, str]) -> str | None:
    if not header:
        return None
    if header in aliases:
        return aliases[header]
    for alias, key in aliases.items():
        if alias and (header.startswith(alias) or alias.startswith(header)):
            return key
    return None


def list_excel_files(folder: Path) -> list[dict[str, Any]]:
    results = []
    for path in sorted(folder.glob("*.xls*")):
        if path.name.startswith("~$"):
            continue
        stat = path.stat()
        results.append({
            "name": path.name,
            "size": stat.st_size,
            "mtime": stat.st_mtime,
        })
    return results


def sheet_names(path: Path) -> list[str]:
    if not HAS_OPENPYXL:
        raise RuntimeError("openpyxl-missing")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


# ---------------------------------------------------------------- import
def read_task_sheet(path: Path, sheet: str, type_key: str) -> list[dict[str, Any]]:
    """Read a task sheet, auto-detecting the header row."""
    if not HAS_OPENPYXL:
        raise RuntimeError("openpyxl-missing")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in workbook.sheetnames:
            return []
        worksheet = workbook[sheet]
        aliases = _alias_map(type_key)
        columns: dict[int, str] = {}
        rows: list[dict[str, Any]] = []
        blank_run = 0
        header_found = False

        for row_index, row in enumerate(worksheet.iter_rows(values_only=True), 1):
            if not header_found:
                if row_index > 15:
                    break
                headers = [model.norm_header(c) for c in row]
                if KEY_HEADER not in headers:
                    continue
                for col_index, header in enumerate(headers):
                    key = _match_key(header, aliases)
                    if key and key not in columns.values():
                        columns[col_index] = key
                header_found = True
                continue

            values = list(row)
            if not any(v not in (None, "") for v in values):
                blank_run += 1
                if blank_run >= MAX_BLANK_RUN:
                    break
                continue
            blank_run = 0
            if values and model.is_sample_row(values[0]):
                continue

            record: dict[str, Any] = {"type": type_key}
            for col_index, key in columns.items():
                if col_index < len(values):
                    record[key] = values[col_index]
            if not (record.get("programId") or record.get("project")):
                continue
            rows.append(record)
        return rows
    finally:
        workbook.close()


def read_entry_sheet(path: Path, sheet: str) -> dict[str, dict[str, str]]:
    """Read the monthly '実績入力確認' matrix -> {member: {'YYYY-MM': status}}."""
    if not HAS_OPENPYXL:
        raise RuntimeError("openpyxl-missing")
    name_header = model.norm_header("氏名")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in workbook.sheetnames:
            return {}
        worksheet = workbook[sheet]
        grid = [list(r) for r in worksheet.iter_rows(min_row=1, max_row=200, values_only=True)]
        name_col = None
        header_row = None
        year = None
        for r_index, row in enumerate(grid):
            for c_index, cell in enumerate(row):
                if model.norm_header(cell) == name_header:
                    name_col, header_row = c_index, r_index
                    for other in row:
                        digits = "".join(ch for ch in model.norm_header(other) if ch.isdigit())
                        if len(digits) == 4:
                            year = int(digits)
                    break
            if name_col is not None:
                break
        if name_col is None or header_row is None:
            return {}

        month_cols: dict[int, int] = {}
        for row in grid[header_row: header_row + 3]:
            for c_index, cell in enumerate(row):
                digits = "".join(ch for ch in model.norm_header(cell) if ch.isdigit())
                if digits and c_index > name_col and 1 <= int(digits) <= 12 and len(digits) <= 2:
                    month_cols[c_index] = int(digits)
            if month_cols:
                break
        if not month_cols:
            return {}

        result: dict[str, dict[str, str]] = {}
        for row in grid[header_row + 1:]:
            if name_col >= len(row):
                continue
            member = model.norm_member(row[name_col])
            if not member or member == "氏名":
                continue
            months: dict[str, str] = {}
            for c_index, month in month_cols.items():
                if c_index >= len(row):
                    continue
                value = row[c_index]
                if value in (None, ""):
                    continue
                cal_year = year if year is not None else 0
                if cal_year and month < 4:
                    cal_year += 1
                months[f"{cal_year:04d}-{month:02d}"] = str(value).strip()
            if months:
                result[member] = months
        return result
    finally:
        workbook.close()


def import_workbook(path: Path) -> dict[str, Any]:
    """Import every known sheet from a workbook."""
    names = sheet_names(path)
    tasks: list[dict[str, Any]] = []
    detail: dict[str, int] = {}
    for type_def in model.SCHEMA["types"]:
        sheet = _find_sheet(names, type_def["sheet"])
        if not sheet:
            continue
        rows = read_task_sheet(path, sheet, type_def["key"])
        detail[type_def["label"]] = len(rows)
        tasks.extend(rows)
    entries: dict[str, dict[str, str]] = {}
    entry_sheet = _find_sheet(names, "実績入力確認")
    if entry_sheet:
        entries = read_entry_sheet(path, entry_sheet)
    return {"tasks": tasks, "entries": entries, "detail": detail, "sheets": names}


def _find_sheet(names: Iterable[str], wanted: str) -> str | None:
    target = model.norm_header(wanted)
    for name in names:
        if model.norm_header(name) == target:
            return name
    for name in names:
        if target in model.norm_header(name):
            return name
    return None


# ------------------------------------------------------------------- csv
def read_csv(raw: bytes, type_key: str) -> list[dict[str, Any]]:
    text = _decode(raw)
    reader = csv.DictReader(io.StringIO(text))
    aliases = _alias_map(type_key)
    rows: list[dict[str, Any]] = []
    for line in reader:
        record: dict[str, Any] = {"type": type_key}
        for header, value in line.items():
            key = _match_key(model.norm_header(header), aliases)
            if key and value not in (None, ""):
                record[key] = value
        if record.get("programId") or record.get("project"):
            rows.append(record)
    return rows


def _decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp932"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def write_csv(tasks: list[dict], type_key: str, keys: list[str] | None = None) -> bytes:
    fields = [f for f in model.SCHEMA["fields"] if type_key in f["types"]]
    if keys:
        order = {k: i for i, k in enumerate(keys)}
        fields = [f for f in fields if f["key"] in order]
        fields.sort(key=lambda f: order[f["key"]])
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow([f["label"] for f in fields])
    for task in tasks:
        writer.writerow(["" if task.get(f["key"]) is None else task.get(f["key"]) for f in fields])
    return buffer.getvalue().encode("utf-8-sig")


# ----------------------------------------------------------------- excel
HEADER_FILL = "FF1F3864"


def write_workbook(tasks: list[dict], entries: dict, masters: dict) -> bytes:
    if not HAS_OPENPYXL:
        raise RuntimeError("openpyxl-missing")
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    for type_def in model.SCHEMA["types"]:
        type_key = type_def["key"]
        fields = [f for f in model.SCHEMA["fields"] if type_key in f["types"]]
        sheet = workbook.create_sheet(type_def["sheet"])
        sheet.append([f["label"] for f in fields])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFFFF")
            cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for task in tasks:
            if task.get("type") != type_key:
                continue
            sheet.append([task.get(f["key"]) for f in fields])
        for index, field in enumerate(fields, 1):
            sheet.column_dimensions[get_column_letter(index)].width = max(
                10, min(46, int(field.get("width", 100) / 7))
            )
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions

    if entries:
        sheet = workbook.create_sheet("実績入力確認")
        months = sorted({m for row in entries.values() for m in row})
        sheet.append(["氏名"] + months)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFFFF")
            cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        for member, row in entries.items():
            sheet.append([member] + [row.get(m, "") for m in months])
        sheet.freeze_panes = "B2"

    sheet = workbook.create_sheet("マスタ")
    sheet.append(["区分", "値"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
    for key in ("members", "projects", "resources", "systems", "kubuns", "editors"):
        for value in masters.get(key) or []:
            sheet.append([key, value])
    for key, value in (masters.get("settings") or {}).items():
        sheet.append([f"settings.{key}", value])

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()
